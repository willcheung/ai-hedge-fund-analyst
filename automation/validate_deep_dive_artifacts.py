#!/usr/bin/env python3
"""Validate wiki-market deep-dive model/chart artifacts.

Usage:
  python3 ${ANALYST_HERMES_HOME}/scripts/validate_deep_dive_artifacts.py TICKER YYYY-MM-DD

Checks the artifact contract used by the cross-checked deep-dive workflow:
- ${ANALYST_WIKI_ROOT}/models/{TICKER}*{DATE}*.xlsx exists
- ${ANALYST_WIKI_ROOT}/charts/{TICKER}*{DATE}*.png exists unless --no-chart
- Workbook has Checks and Methodology/Assumptions-style tabs
- Workbook contains formulas and hardcoded input comments
- Checks tab has formulas, so Excel can surface TRUE/FALSE gates
- Workbook includes a forward-looking signal layer: guidance/consensus/backlog/projection evidence
"""
from __future__ import annotations
from automation_paths import configured_text

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

WIKI = Path(configured_text('${ANALYST_WIKI_ROOT}'))


def fail(msg: str) -> None:
    print(f"FAIL: {msg}")
    sys.exit(1)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('ticker')
    ap.add_argument('date')
    ap.add_argument('--no-chart', action='store_true', help='Allow no PNG chart when chart quality gate failed')
    args = ap.parse_args()

    ticker = args.ticker.upper().lstrip('$')
    date = args.date
    models = sorted((WIKI / 'models').glob(f'{ticker}*{date}*.xlsx'))
    if not models:
        fail(configured_text(f'No model found matching ${{ANALYST_WIKI_ROOT}}/models/{ticker}*{date}*.xlsx'))
    model = models[-1]

    charts = sorted((WIKI / 'charts').glob(f'{ticker}*{date}*.png'))
    if not charts and not args.no_chart:
        fail(configured_text(f'No chart found matching ${{ANALYST_WIKI_ROOT}}/charts/{ticker}*{date}*.png'))

    wb = load_workbook(model, data_only=False)
    sheet_names = set(wb.sheetnames)

    if 'Checks' not in sheet_names:
        fail('Workbook missing Checks tab')

    has_methodology = any(s in sheet_names for s in ('Methodology', 'Assumptions', 'Inputs'))
    if not has_methodology:
        fail('Workbook missing Methodology/Assumptions/Inputs tab')

    formula_cells = []
    commented_inputs = []
    blue_inputs = []
    checks_formulas = []
    workbook_text = []

    for ws in wb.worksheets:
        workbook_text.append(ws.title)
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str):
                    workbook_text.append(val)
                if isinstance(val, str) and val.startswith('='):
                    formula_cells.append(f'{ws.title}!{cell.coordinate}')
                    if ws.title == 'Checks':
                        checks_formulas.append(f'{ws.title}!{cell.coordinate}')
                if val is not None and cell.comment is not None:
                    commented_inputs.append(f'{ws.title}!{cell.coordinate}')
                    workbook_text.append(str(cell.comment.text))
                color = getattr(cell.font, 'color', None)
                rgb = getattr(color, 'rgb', None) if color else None
                if rgb and str(rgb).upper().endswith('0000FF'):
                    blue_inputs.append(f'{ws.title}!{cell.coordinate}')

    text_blob = ' '.join(workbook_text).lower()
    forward_hits = [
        kw for kw in (
            'forward signal', 'guidance', 'consensus', 'forecast', 'projected',
            'fy+1', 'fy+2', 'backlog', 'orders', 'bookings', 'rpo',
            'earnings call', 'investor presentation', 'capacity', 'customer ramp'
        ) if kw in text_blob
    ]

    if len(formula_cells) < 5:
        fail(f'Workbook has too few formulas ({len(formula_cells)}); model should be formula-driven')
    if not checks_formulas:
        fail('Checks tab has no formulas')
    if len(commented_inputs) < 5:
        fail(f'Workbook has too few sourced/commented inputs ({len(commented_inputs)})')
    if len(blue_inputs) < 5:
        fail(f'Workbook has too few blue assumption/input cells ({len(blue_inputs)})')
    if len(forward_hits) < 2:
        fail('Workbook missing forward-looking signal layer: add guidance/consensus/backlog/projection inputs with sources')

    print('PASS')
    print(f'Model: {model}')
    if charts:
        print(f'Chart: {charts[-1]}')
    else:
        print('Chart: skipped by --no-chart')
    print(f'Sheets: {", ".join(wb.sheetnames)}')
    print(f'Formulas: {len(formula_cells)} | Commented inputs: {len(commented_inputs)} | Blue inputs: {len(blue_inputs)} | Check formulas: {len(checks_formulas)} | Forward hits: {", ".join(forward_hits)}')


if __name__ == '__main__':
    main()
